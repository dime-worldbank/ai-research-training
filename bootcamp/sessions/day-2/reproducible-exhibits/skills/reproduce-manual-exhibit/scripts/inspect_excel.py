#!/usr/bin/env python3
"""Inventory workbook structure, formulas, styles, tables, and chart references."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any


def _title_text(chart: Any) -> str | None:
    try:
        paragraphs = chart.title.tx.rich.p
        return " ".join(run.t for p in paragraphs for run in p.r if run.t).strip() or None
    except (AttributeError, TypeError):
        return None


def _series_formula(series: Any, attr: str) -> str | None:
    obj = getattr(series, attr, None)
    if obj is None:
        return None
    for path in (("numRef", "f"), ("strRef", "f")):
        current = obj
        try:
            for name in path:
                current = getattr(current, name)
            if current:
                return str(current)
        except AttributeError:
            pass
    return None


def inspect_workbook(path: Path, formula_limit: int) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("openpyxl is required: install it in the active Python environment") from exc

    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    cached = openpyxl.load_workbook(path, data_only=True, read_only=False)
    report: dict[str, Any] = {
        "file": str(path.resolve()),
        "sheet_count": len(workbook.sheetnames),
        "defined_names": sorted(str(name) for name in workbook.defined_names),
        "sheets": [],
    }

    for sheet in workbook.worksheets:
        cached_sheet = cached[sheet.title]
        formulas = []
        formula_count = 0
        styled_nonempty = 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.style_id:
                    styled_nonempty += 1
                if cell.data_type == "f":
                    formula_count += 1
                    if len(formulas) < formula_limit:
                        formulas.append(
                            {
                                "cell": cell.coordinate,
                                "formula": cell.value,
                                "cached_value": cached_sheet[cell.coordinate].value,
                            }
                        )

        charts = []
        for chart in sheet._charts:
            series = [
                {
                    "title": _series_formula(item, "tx"),
                    "categories": _series_formula(item, "cat") or _series_formula(item, "xVal"),
                    "values": _series_formula(item, "val") or _series_formula(item, "yVal"),
                }
                for item in chart.ser
            ]
            anchor = getattr(chart, "anchor", None)
            anchor_from = getattr(anchor, "_from", None)
            charts.append(
                {
                    "type": type(chart).__name__,
                    "title": _title_text(chart),
                    "anchor": (
                        {"column": anchor_from.col + 1, "row": anchor_from.row + 1}
                        if anchor_from is not None
                        else None
                    ),
                    "series": series,
                }
            )

        report["sheets"].append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "dimensions": sheet.calculate_dimension(),
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                "tables": sorted(sheet.tables.keys()),
                "formula_count": formula_count,
                "formulas_truncated": formula_count > len(formulas),
                "formulas": formulas,
                "styled_nonempty_cells": styled_nonempty,
                "charts": charts,
            }
        )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, help="Write JSON to this path instead of stdout")
    parser.add_argument("--formula-limit", type=int, default=200)
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise SystemExit(f"Workbook not found: {args.workbook}")
    report = inspect_workbook(args.workbook, max(args.formula_limit, 0))
    encoded = json.dumps(report, indent=2, ensure_ascii=False, default=str) + "\n"
    if args.output:
        args.output.write_text(encoded, encoding="utf-8")
    else:
        print(encoded, end="")


if __name__ == "__main__":
    main()
